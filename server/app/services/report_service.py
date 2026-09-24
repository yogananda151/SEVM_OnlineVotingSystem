import io
from datetime import datetime
from typing import Optional, Dict, Any
from sqlalchemy.orm import Session
from sqlalchemy import func
from fastapi import HTTPException
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from app.models.election import Election, ElectionConstituency
from app.models.party_candidate import Candidate
from app.models.voter import Voter, ElectionVoterStatus
from app.models.vote import Vote
from app.models.audit import AuditLog
from app.models.location import PollingStation

class ReportService:
    @staticmethod
    def generate_election_summary_pdf(db: Session, election_id: int) -> io.BytesIO:
        election = db.query(Election).filter(Election.id == election_id).first()
        if not election:
            raise HTTPException(status_code=404, detail="Election not found.")

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "TitleStyle",
            parent=styles["Heading1"],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#1A73E8"),
            alignment=1,
        )
        subtitle_style = ParagraphStyle(
            "SubTitleStyle",
            parent=styles["Normal"],
            fontSize=13,
            leading=16,
            textColor=colors.HexColor("#333333"),
            alignment=1,
        )
        meta_style = ParagraphStyle(
            "MetaStyle",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#666666"),
            alignment=1,
        )
        heading_style = ParagraphStyle(
            "HeadingStyle",
            parent=styles["Heading2"],
            fontSize=12,
            leading=15,
            textColor=colors.HexColor("#1A73E8"),
        )
        body_style = ParagraphStyle(
            "BodyStyle",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#333333"),
        )

        elements = [
            Paragraph("Election Commission of India", title_style),
            Paragraph("ELECTION SUMMARY REPORT", subtitle_style),
            Paragraph(f"Generated: {datetime.now().strftime('%d/%m/%Y, %H:%M:%S')}", meta_style),
            Spacer(1, 15),
            Paragraph("Election Details", heading_style),
            Spacer(1, 4),
        ]

        info_data = [
            [Paragraph("<b>Election Name:</b>", body_style), Paragraph(election.name, body_style)],
            [Paragraph("<b>Type:</b>", body_style), Paragraph(election.electionType, body_style)],
            [Paragraph("<b>Status:</b>", body_style), Paragraph(str(election.status.value if hasattr(election.status, "value") else election.status), body_style)],
            [Paragraph("<b>Scheduled Date:</b>", body_style), Paragraph(election.scheduledDate.strftime("%d/%m/%Y"), body_style)],
        ]
        if election.startTime:
            info_data.append([Paragraph("<b>Started:</b>", body_style), Paragraph(election.startTime.strftime("%d/%m/%Y, %H:%M"), body_style)])
        if election.endTime:
            info_data.append([Paragraph("<b>Ended:</b>", body_style), Paragraph(election.endTime.strftime("%d/%m/%Y, %H:%M"), body_style)])

        info_table = Table(info_data, colWidths=[120, 395])
        info_table.setStyle(TableStyle([
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#333333")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
        ]))
        elements.append(info_table)
        elements.append(Spacer(1, 15))

        # Constituencies & candidates
        links = db.query(ElectionConstituency).filter(ElectionConstituency.electionId == election_id).all()
        for link in links:
            con = link.constituency
            elements.append(Paragraph(f"Constituency: {con.name} ({con.code})", heading_style))
            elements.append(Spacer(1, 4))

            cands = (
                db.query(Candidate)
                .filter(Candidate.electionId == election_id, Candidate.constituencyId == con.id, Candidate.deletedAt.is_(None))
                .order_by(Candidate.serialNumber.asc())
                .all()
            )

            table_data = [["#", "Candidate Name", "Party", "Votes"]]
            for c in cands:
                vote_count = db.query(func.count(Vote.id)).filter(Vote.candidateId == c.id, Vote.electionId == election_id).scalar() or 0
                party_name = c.party.name if c.party else "Independent"
                table_data.append([str(c.serialNumber), c.fullName, party_name, str(vote_count)])

            cand_table = Table(table_data, colWidths=[30, 200, 200, 85])
            cand_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A73E8")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E0E0E0")),
            ]))
            elements.append(cand_table)
            elements.append(Spacer(1, 15))

        elements.append(Spacer(1, 10))
        elements.append(Paragraph("OFFICIAL DOCUMENT – ELECTION COMMISSION OF INDIA", meta_style))

        doc.build(elements)
        buffer.seek(0)
        return buffer

    @staticmethod
    def generate_audit_log_pdf(db: Session, start_date: Optional[datetime] = None, end_date: Optional[datetime] = None) -> io.BytesIO:
        query = db.query(AuditLog)
        if start_date:
            query = query.filter(AuditLog.createdAt >= start_date)
        if end_date:
            query = query.filter(AuditLog.createdAt <= end_date)
        logs = query.order_by(AuditLog.createdAt.desc()).limit(500).all()

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=35, leftMargin=35, topMargin=35, bottomMargin=35)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "TitleStyle",
            parent=styles["Heading1"],
            fontSize=16,
            textColor=colors.HexColor("#1A73E8"),
            alignment=1,
        )
        meta_style = ParagraphStyle(
            "MetaStyle",
            parent=styles["Normal"],
            fontSize=8,
            textColor=colors.HexColor("#666666"),
            alignment=1,
        )
        log_style = ParagraphStyle(
            "LogStyle",
            parent=styles["Normal"],
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#333333"),
        )

        elements = [
            Paragraph("AUDIT LOG REPORT", title_style),
            Paragraph(f"Generated: {datetime.now().strftime('%d/%m/%Y, %H:%M:%S')}", meta_style),
            Spacer(1, 15),
        ]

        table_data = [["Timestamp", "Action", "Module", "User", "Description"]]
        for log in logs:
            user_str = log.user.email if log.user else "System"
            action_str = str(log.action.value if hasattr(log.action, "value") else log.action)
            table_data.append([
                log.createdAt.strftime("%d/%m %H:%M"),
                action_str,
                log.module,
                user_str,
                Paragraph(log.description[:120], log_style),
            ])

        log_table = Table(table_data, colWidths=[70, 70, 65, 110, 210])
        log_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A73E8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#EEEEEE")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        elements.append(log_table)

        doc.build(elements)
        buffer.seek(0)
        return buffer

    @staticmethod
    def generate_results_excel(db: Session, election_id: int) -> io.BytesIO:
        election = db.query(Election).filter(Election.id == election_id).first()
        if not election:
            raise HTTPException(status_code=404, detail="Election not found.")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Election Results"

        # Title
        ws.merge_cells("A1:F1")
        ws["A1"] = f"Election Results – {election.name}"
        ws["A1"].font = Font(name="Arial", size=14, bold=True, color="1A73E8")
        ws["A1"].alignment = Alignment(horizontal="center")

        ws.merge_cells("A2:F2")
        ws["A2"] = f"Generated: {datetime.now().strftime('%d/%m/%Y, %H:%M:%S')}"
        ws["A2"].alignment = Alignment(horizontal="center")

        row = 4
        links = db.query(ElectionConstituency).filter(ElectionConstituency.electionId == election_id).all()
        for link in links:
            con = link.constituency
            ws.cell(row=row, column=1, value=f"Constituency: {con.name} ({con.code})").font = Font(name="Arial", size=12, bold=True, color="1A73E8")
            row += 1

            headers = ["#", "Candidate Name", "Party", "Abbreviation", "Votes", "Winner"]
            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=row, column=col_idx, value=h)
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="1A73E8", end_color="1A73E8", fill_type="solid")
            row += 1

            cands = (
                db.query(Candidate)
                .filter(Candidate.electionId == election_id, Candidate.constituencyId == con.id, Candidate.deletedAt.is_(None))
                .all()
            )
            cand_rows = []
            for c in cands:
                vote_count = db.query(func.count(Vote.id)).filter(Vote.candidateId == c.id, Vote.electionId == election_id).scalar() or 0
                cand_rows.append((c, vote_count))
            cand_rows.sort(key=lambda x: x[1], reverse=True)

            max_votes = cand_rows[0][1] if cand_rows else 0

            for c, vote_count in cand_rows:
                is_winner = election.isResultPublished and vote_count == max_votes and max_votes > 0
                ws.cell(row=row, column=1, value=c.serialNumber)
                ws.cell(row=row, column=2, value=c.fullName)
                ws.cell(row=row, column=3, value=c.party.name if c.party else "Independent")
                ws.cell(row=row, column=4, value=c.party.abbreviation if c.party else "IND")
                ws.cell(row=row, column=5, value=vote_count if election.isResultPublished else "Locked")
                ws.cell(row=row, column=6, value="🏆 Winner" if is_winner else "")

                if is_winner:
                    for col_idx in range(1, 7):
                        ws.cell(row=row, column=col_idx).fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
                row += 1
            row += 1

        widths = [8, 30, 35, 15, 12, 15]
        for idx, w in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = w

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    @staticmethod
    def generate_voters_excel(db: Session, station_id: int) -> io.BytesIO:
        voters = (
            db.query(Voter)
            .filter(Voter.pollingStationId == station_id, Voter.deletedAt.is_(None))
            .order_by(Voter.serialNumber.asc())
            .all()
        )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Voters List"

        headers = ["#", "Full Name", "Voter ID", "Date of Birth", "Gender", "Address", "Has Voted", "Voted At"]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="1A73E8", end_color="1A73E8", fill_type="solid")

        active_election = db.query(Election).filter(Election.status == "ACTIVE").first()

        for idx, v in enumerate(voters, start=2):
            has_voted = False
            voted_at = "-"
            if active_election:
                status_rec = (
                    db.query(ElectionVoterStatus)
                    .filter(ElectionVoterStatus.voterId == v.id, ElectionVoterStatus.electionId == active_election.id)
                    .first()
                )
                if status_rec and status_rec.hasVoted:
                    has_voted = True
                    voted_at = status_rec.votedAt.strftime("%d/%m/%Y, %H:%M") if status_rec.votedAt else "-"

            ws.cell(row=idx, column=1, value=v.serialNumber)
            ws.cell(row=idx, column=2, value=v.fullName)
            ws.cell(row=idx, column=3, value=v.voterId)
            ws.cell(row=idx, column=4, value=v.dateOfBirth.strftime("%d/%m/%Y"))
            ws.cell(row=idx, column=5, value=v.gender)
            ws.cell(row=idx, column=6, value=v.address)
            ws.cell(row=idx, column=7, value="Yes" if has_voted else "No")
            ws.cell(row=idx, column=8, value=voted_at)

            if has_voted:
                ws.cell(row=idx, column=7).fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")

        widths = [8, 30, 20, 15, 10, 40, 12, 20]
        for col_idx, w in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = w

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

report_service = ReportService()
