import io
from typing import Dict, Any, List, Optional
from sqlalchemy.orm import Session
from fastapi import HTTPException
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from app.models.location import Constituency
from app.models.party_candidate import PoliticalParty, Candidate

class CandidateExcelService:
    @staticmethod
    def generate_template(db: Session) -> io.BytesIO:
        constituencies = db.query(Constituency).filter(Constituency.deletedAt.is_(None)).order_by(Constituency.id.asc()).all()
        parties = db.query(PoliticalParty).filter(PoliticalParty.deletedAt.is_(None)).order_by(PoliticalParty.id.asc()).all()

        wb = openpyxl.Workbook()
        ws_roster = wb.active
        ws_roster.title = "Candidates"

        headers = [
            "Full Name *", "Age *", "Serial Number", "Qualification",
            "Independent (Yes/No) *", "Party ID (If not Independent)", "Constituency ID *",
        ]
        ws_roster.append(headers)

        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")

        for col_num in range(1, len(headers) + 1):
            cell = ws_roster.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        sample_c_id = constituencies[0].id if constituencies else 1
        sample_p_id = parties[0].id if parties else 1

        samples = [
            ["Jane Doe", 45, 1, "B.A.", "No", sample_p_id, sample_c_id],
            ["John Smith", 50, 2, "M.Sc.", "Yes", "", sample_c_id],
        ]
        for row in samples:
            ws_roster.append(row)

        widths = [26, 10, 15, 25, 22, 26, 18]
        for idx, width in enumerate(widths, start=1):
            col_letter = openpyxl.utils.get_column_letter(idx)
            ws_roster.column_dimensions[col_letter].width = width

        # Reference sheet
        ws_ref = wb.create_sheet(title="Reference Data")
        ref_headers = ["Constituency ID", "Constituency Name", "", "Party ID", "Party Name", "Party Abbr"]
        ws_ref.append(ref_headers)

        ref_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        for col_num in range(1, len(ref_headers) + 1):
            if col_num == 3:
                continue
            cell = ws_ref.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = ref_fill
            cell.alignment = header_align

        max_rows = max(len(constituencies), len(parties))
        for i in range(max_rows):
            c = constituencies[i] if i < len(constituencies) else None
            p = parties[i] if i < len(parties) else None
            ws_ref.append([
                c.id if c else "",
                c.name if c else "",
                "",
                p.id if p else "",
                p.name if p else "",
                p.abbreviation if p else "",
            ])

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    @staticmethod
    def parse_excel(
        db: Session,
        file_bytes: bytes,
        election_id: int,
        default_constituency_id: Optional[int] = None,
    ) -> Dict[str, Any]:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            ws = wb.worksheets[0]
        except Exception:
            raise HTTPException(status_code=400, detail="Failed to parse spreadsheet. Please upload a valid .xlsx file.")

        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            raise HTTPException(status_code=400, detail="The uploaded sheet is empty or contains no data rows.")

        constituencies = db.query(Constituency).filter(Constituency.deletedAt.is_(None)).all()
        parties = db.query(PoliticalParty).filter(PoliticalParty.deletedAt.is_(None)).all()

        con_by_id = {c.id: c for c in constituencies}
        con_by_name = {c.name.strip().lower(): c for c in constituencies}
        con_by_code = {c.code.strip().lower(): c for c in constituencies}

        party_by_id = {p.id: p for p in parties}
        party_by_abbr = {p.abbreviation.strip().lower(): p for p in parties}
        party_by_name = {p.name.strip().lower(): p for p in parties}

        header_row = [str(cell or "").lower().strip() for cell in rows[0]]
        col_map = {}
        for idx, h in enumerate(header_row):
            h_clean = "".join(ch for ch in h if ch.isalnum())
            if "partyid" in h_clean:
                col_map["partyId"] = idx
            elif "independent" in h_clean and "party" not in h_clean:
                col_map["isIndependent"] = idx
            elif "constituencyid" in h_clean:
                col_map["constituencyId"] = idx
            elif any(k in h_clean for k in ["constituencyname", "constituency"]) and "constituencyId" not in col_map:
                col_map["constituency"] = idx
            elif any(k in h_clean for k in ["partyname", "partyabbr", "party"]) and "partyId" not in col_map:
                col_map["party"] = idx
            elif any(k in h_clean for k in ["fullname", "candidatename", "name"]) and "party" not in h_clean and "constituency" not in h_clean:
                col_map["fullName"] = idx
            elif "age" in h_clean:
                col_map["age"] = idx
            elif any(k in h_clean for k in ["serialnumber", "serialno", "serial", "ballot"]):
                col_map["serialNumber"] = idx
            elif any(k in h_clean for k in ["qualification", "education"]):
                col_map["qualification"] = idx

        if "fullName" not in col_map:
            raise HTTPException(status_code=400, detail='Could not find "Full Name" column.')

        errors = []
        duplicate_names = []
        seen_names = set()
        raw_candidates = []

        for r_idx, row in enumerate(rows[1:], start=2):
            if not any(row):
                continue

            full_name = str(row[col_map["fullName"]] or "").strip()
            if not full_name or len(full_name) < 2:
                errors.append(f"Row {r_idx}: Full Name is required.")
                continue

            # Resolve Constituency
            con = None
            if "constituencyId" in col_map and row[col_map["constituencyId"]] is not None:
                val = str(row[col_map["constituencyId"]]).strip()
                if val.isdigit():
                    con = con_by_id.get(int(val))
                if not con:
                    con = con_by_name.get(val.lower()) or con_by_code.get(val.lower())

            if not con and "constituency" in col_map and row[col_map["constituency"]] is not None:
                val = str(row[col_map["constituency"]]).strip()
                if val.isdigit():
                    con = con_by_id.get(int(val))
                if not con:
                    con = con_by_name.get(val.lower()) or con_by_code.get(val.lower())

            if not con and default_constituency_id:
                con = con_by_id.get(default_constituency_id)

            if not con:
                errors.append(f"Row {r_idx}: Valid Constituency ID or Name is required.")
                continue
            constituency_id = con.id

            dedup_key = f"{full_name.lower()}-{constituency_id}"
            if dedup_key in seen_names:
                duplicate_names.append(full_name)
                continue
            seen_names.add(dedup_key)

            raw_age = row[col_map.get("age")] if "age" in col_map else None
            try:
                age = int(float(str(raw_age).strip())) if raw_age is not None and str(raw_age).strip() else 35
            except Exception:
                age = 35

            raw_serial = row[col_map.get("serialNumber")] if "serialNumber" in col_map else None
            try:
                serial_no = int(float(str(raw_serial).strip())) if raw_serial is not None and str(raw_serial).strip() else len(raw_candidates) + 1
            except Exception:
                serial_no = len(raw_candidates) + 1

            indep_str = str(row[col_map.get("isIndependent")] or "").lower().strip() if "isIndependent" in col_map else "no"
            is_independent = indep_str.startswith("y") or indep_str == "true" or indep_str == "1"

            # Resolve Party
            party_id = None
            if not is_independent:
                p_val = None
                if "partyId" in col_map and row[col_map["partyId"]] is not None:
                    p_val = str(row[col_map["partyId"]]).strip()
                elif "party" in col_map and row[col_map["party"]] is not None:
                    p_val = str(row[col_map["party"]]).strip()

                if p_val:
                    if p_val.isdigit():
                        p_obj = party_by_id.get(int(p_val))
                    else:
                        p_obj = party_by_abbr.get(p_val.lower()) or party_by_name.get(p_val.lower())
                    if p_obj:
                        party_id = p_obj.id
                    else:
                        errors.append(f"Row {r_idx}: Political Party '{p_val}' not found.")
                        continue
                else:
                    is_independent = True

            qualification = str(row[col_map.get("qualification")] or "").strip() if "qualification" in col_map else None

            raw_candidates.append({
                "electionId": election_id,
                "constituencyId": constituency_id,
                "fullName": full_name,
                "age": age,
                "serialNumber": serial_no,
                "qualification": qualification,
                "isIndependent": is_independent,
                "partyId": party_id,
            })

        # Check existing in DB
        existing = db.query(Candidate.fullName, Candidate.constituencyId, Candidate.partyId, Candidate.isIndependent).filter(
            Candidate.electionId == election_id,
            Candidate.deletedAt.is_(None),
        ).all()
        existing_set = {f"{c[0].lower()}-{c[1]}" for c in existing}
        existing_parties_in_con = {f"{c[1]}-{c[2]}" for c in existing if c[2] and not c[3]}

        seen_parties_in_batch = set()
        valid_candidates = []
        for c in raw_candidates:
            if f"{c['fullName'].lower()}-{c['constituencyId']}" in existing_set:
                duplicate_names.append(c["fullName"])
                continue

            if c["partyId"] and not c["isIndependent"]:
                p_key = f"{c['constituencyId']}-{c['partyId']}"
                if p_key in existing_parties_in_con or p_key in seen_parties_in_batch:
                    errors.append(f"Candidate '{c['fullName']}': A candidate from this party is already registered in constituency ID {c['constituencyId']}. Max 1 candidate per party allowed.")
                    continue
                seen_parties_in_batch.add(p_key)

            valid_candidates.append(c)

        return {
            "totalRows": len(rows) - 1,
            "validRowsCount": len(raw_candidates),
            "importedCount": len(valid_candidates),
            "skippedDuplicatesCount": len(duplicate_names),
            "duplicateCandidateNames": duplicate_names,
            "errors": errors,
            "validCandidates": valid_candidates,
        }

candidate_excel_service = CandidateExcelService()
